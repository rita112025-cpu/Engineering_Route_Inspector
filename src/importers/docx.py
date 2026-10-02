"""DOCX specification import (python-docx).

DOCX files do not store page numbers (pagination happens at render time),
so ``page`` is None and line numbers are paragraph numbers.
"""
from __future__ import annotations

from pathlib import Path

from core.evidence.chunks import chunk_lines, evidence_id, sha256_text
from core.models.results import EvidenceChunk
from .text import DocumentImportError


def load_docx(path: str | Path, document_id: str, doc_hash: str, filename: str | None = None) -> list[EvidenceChunk]:
    try:
        import docx
    except ImportError as exc:
        raise DocumentImportError("無法讀取 DOCX", "缺少 python-docx 套件，請重新執行 start-ui.bat 安裝。", repr(exc))
    path = Path(path)
    fname = filename or path.name
    try:
        d = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise DocumentImportError("無法讀取 DOCX", "檔案損毀或不是 Word (.docx) 格式。", repr(exc))
    chunks: list[EvidenceChunk] = []
    section = ""
    lineno = 0
    for para in d.paragraphs:
        lineno += 1
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name if para.style is not None else "") or ""
        if style.lower().startswith("heading") or style.startswith("標題"):
            section = text
        c, section = chunk_lines([text], document_id=document_id, filename=fname, doc_hash=doc_hash,
                                 page=None, section=section, line_offset=lineno - 1)
        chunks.extend(c)
    for ti, table in enumerate(d.tables, start=1):
        for ri, row in enumerate(table.rows, start=1):
            cells = [c.text.strip() for c in row.cells]
            text = " | ".join(dict.fromkeys(c for c in cells if c))
            if not text:
                continue
            ln = 100000 * ti + ri
            chunks.append(EvidenceChunk(
                id=evidence_id(doc_hash, None, ln, text), document_id=document_id, filename=fname,
                page=None, line_start=ln, line_end=ln, section=f"表格 {ti} 第 {ri} 列", text=text,
                hash=sha256_text(text)))
    return chunks
