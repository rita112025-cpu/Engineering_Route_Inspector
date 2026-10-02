import math

import ezdxf
import pytest

from importers.docx import load_docx
from importers.dxf import DrawingImportError, load_dxf
from importers.pdf import load_pdf
from importers.text import DocumentImportError, decode_bytes, load_text


def _save(doc, tmp_path, name="plan.dxf"):
    p = tmp_path / name
    doc.saveas(p)
    return p


def test_basic_entities_and_units(tmp_path):
    doc = ezdxf.new("R2018")
    doc.header["$INSUNITS"] = 6  # metres
    msp = doc.modelspace()
    msp.add_line((0, 0), (10, 0), dxfattribs={"layer": "SCADA"})
    msp.add_lwpolyline([(0, 0), (5, 0), (5, 5)], close=True, dxfattribs={"layer": "ZONE"})
    msp.add_circle((3, 3), 1.5, dxfattribs={"layer": "SUPPORT"})
    msp.add_arc((0, 0), 2, 0, 90, dxfattribs={"layer": "ARC"})
    msp.add_text("SCADA-01", height=0.25, dxfattribs={"layer": "NOTE", "insert": (1, 1)})
    msp.add_mtext("多行\\P文字", dxfattribs={"layer": "NOTE", "insert": (2, 2), "char_height": 0.3})
    msp.add_point((0, 0), dxfattribs={"layer": "PT"})  # unsupported -> skipped, counted
    d = load_dxf(_save(doc, tmp_path))
    assert d.unit_to_mm == 1000.0 and not d.units_assumed and d.units_code == 6
    by_layer = {e.layer: e for e in d.entities}
    assert by_layer["SCADA"].geometry["points"] == [[0.0, 0.0], [10.0, 0.0]]
    assert by_layer["ZONE"].closed and len(by_layer["ZONE"].geometry["points"]) == 3
    assert by_layer["SUPPORT"].kind == "circle" and by_layer["SUPPORT"].geometry["radius"] == 1.5
    arc = by_layer["ARC"]
    assert arc.metadata["source_kind"] == "ARC" and len(arc.geometry["points"]) > 3
    assert arc.length() == pytest.approx(math.pi, rel=1e-3)
    texts = sorted(e.geometry["text"] for e in d.entities if e.kind == "text")
    assert texts == ["SCADA-01", "多行\n文字"]
    assert d.skipped == {"POINT": 1}
    assert d.layers == ["ARC", "NOTE", "SCADA", "SUPPORT", "ZONE"]
    for e in d.entities:
        assert e.id == f"plan.dxf#{e.handle}" and e.source_file == "plan.dxf"


def test_missing_units_are_assumed_mm(tmp_path):
    doc = ezdxf.new("R2018")
    doc.header["$INSUNITS"] = 0
    doc.modelspace().add_line((0, 0), (1, 0))
    d = load_dxf(_save(doc, tmp_path))
    assert d.units_assumed and d.unit_to_mm == 1.0
    assert any("INSUNITS" in w for w in d.warnings)


def test_nested_blocks_transform_layer_inheritance(tmp_path):
    doc = ezdxf.new("R2018")
    inner = doc.blocks.new("INNER")
    inner.add_line((0, 0), (1, 0), dxfattribs={"layer": "0"})
    outer = doc.blocks.new("OUTER")
    outer.add_blockref("INNER", (10, 0), dxfattribs={"layer": "PIPE", "xscale": 2, "yscale": 2})
    outer.add_circle((0, 0), 1, dxfattribs={"layer": "FIXED"})
    msp = doc.modelspace()
    ref = msp.add_blockref("OUTER", (100, 100), dxfattribs={"layer": "ROUTE", "rotation": 90})
    d = load_dxf(_save(doc, tmp_path))
    line = next(e for e in d.entities if e.entity_type == "LINE")
    # INNER line (0,0)-(1,0): scale 2 -> (0,0)-(2,0), +(10,0), rotate 90 -> (0,10)-(0,12), +(100,100)
    (x0, y0), (x1, y1) = line.geometry["points"]
    assert (x0, y0) == pytest.approx((100.0, 110.0)) and (x1, y1) == pytest.approx((100.0, 112.0))
    assert line.layer == "PIPE"                      # layer 0 inherits the INSERT layer
    assert line.metadata["block_path"] == ["OUTER", "INNER"]
    assert line.handle.startswith(ref.dxf.handle + ">")
    circle = next(e for e in d.entities if e.kind == "circle")
    assert circle.layer == "FIXED" and circle.geometry["center"] == pytest.approx([100.0, 100.0])
    assert d.type_counts["INSERT"] == 2


def test_self_referencing_block_is_stopped(tmp_path):
    doc = ezdxf.new("R2018")
    blk = doc.blocks.new("LOOP")
    blk.add_line((0, 0), (1, 0))
    blk.add_blockref("LOOP", (5, 0))
    doc.modelspace().add_blockref("LOOP", (0, 0))
    d = load_dxf(_save(doc, tmp_path))
    assert len([e for e in d.entities if e.entity_type == "LINE"]) == 1
    assert any("遞迴" in w for w in d.warnings)


def test_z_values_recorded(tmp_path):
    doc = ezdxf.new("R2018")
    doc.modelspace().add_line((0, 0, 3000), (10, 0, 3000))
    doc.modelspace().add_line((0, 5), (10, 5))
    d = load_dxf(_save(doc, tmp_path))
    z = sorted((e.metadata["z_known"], e.metadata["z_min"]) for e in d.entities)
    assert z == [(False, 0.0), (True, 3000.0)]


def test_invalid_files_raise_user_errors(tmp_path):
    with pytest.raises(DrawingImportError) as exc:
        load_dxf(tmp_path / "missing.dxf")
    assert exc.value.user_message == "無法讀取 DXF"
    bad = tmp_path / "bad.dxf"
    bad.write_bytes(b"\x00\x01 not a dxf at all \xff")
    with pytest.raises(DrawingImportError):
        load_dxf(bad)


def test_source_file_is_not_modified(tmp_path):
    import hashlib
    doc = ezdxf.new("R2018")
    doc.modelspace().add_line((0, 0), (1, 0))
    p = _save(doc, tmp_path)
    before = hashlib.sha256(p.read_bytes()).hexdigest()
    load_dxf(p)
    assert hashlib.sha256(p.read_bytes()).hexdigest() == before


def test_text_import_encodings(tmp_path):
    p = tmp_path / "spec.txt"
    p.write_bytes("4.1 淨距\n不得小於 300 mm。\n".encode("cp950"))
    chunks = load_text(p, "doc1", "hash")
    assert [c.text for c in chunks] == ["4.1 淨距", "不得小於 300 mm。"]
    assert chunks[1].section == "4.1 淨距" and chunks[1].page == 1 and chunks[1].line_start == 2
    assert decode_bytes("abc".encode("utf-8-sig")) == "abc"
    with pytest.raises(DocumentImportError):
        load_text(tmp_path / "nope.txt", "d", "h")


def test_docx_import(tmp_path):
    import docx
    d = docx.Document()
    d.add_heading("4 管線間距", level=1)
    d.add_paragraph("SCADA 與電力電纜水平淨距不得小於 300 mm。")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "最小淨距"
    t.rows[0].cells[1].text = "300 mm"
    p = tmp_path / "spec.docx"
    d.save(p)
    chunks = load_docx(p, "doc1", "hash")
    body = next(c for c in chunks if "300 mm。" in c.text)
    assert body.section == "4 管線間距" and body.page is None and body.line_start == 2
    table = next(c for c in chunks if c.section.startswith("表格"))
    assert table.text == "最小淨距 | 300 mm"
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a zip")
    with pytest.raises(DocumentImportError):
        load_docx(bad, "d", "h")


def test_pdf_import_pages_and_empty_pages(tmp_path):
    import fitz
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "4.2 Clearance")
    page.insert_text((72, 90), "SCADA cable clearance shall be at least 300 mm.")
    pdf.new_page()  # empty page: no text layer
    page3 = pdf.new_page()
    page3.insert_text((72, 72), "Section 5 no-go zones apply.")
    p = tmp_path / "spec.pdf"
    pdf.save(p)
    chunks, warnings = load_pdf(p, "doc1", "hash")
    assert {c.page for c in chunks} == {1, 3}
    hit = next(c for c in chunks if "300 mm" in c.text)
    assert hit.page == 1 and hit.section == "4.2 Clearance"
    assert warnings and "OCR" in warnings[0]
    bad = tmp_path / "bad.pdf"
    bad.write_bytes(b"%PDF-garbage")
    with pytest.raises(DocumentImportError):
        load_pdf(bad, "d", "h")
